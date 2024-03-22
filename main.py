import argparse
import os
from PyPDF2 import PdfReader
import docx
from openpyxl import load_workbook

# Parse arguments.
parser = argparse.ArgumentParser(description='Extract files metadata.')
parser.add_argument('path', type=str, help='Directory to scan.')
args = parser.parse_args()

# Function to get all files in a directory.
def get_files(path):
    # Scan all files in the directory.
    count = {'pdf': 0, 'docx': 0, 'xlsx': 0}
    files = []
    for current_dir, _, filenames in os.walk(path):
        for filename in filenames:
            fullpath = os.path.join(current_dir, filename)
            # Only append docx, xlsx and pdf files.
            if (fullpath.lower().endswith(('.pdf', '.docx', '.xlsx'))):
                type = os.path.splitext(fullpath)[1].replace('.', '').lower()
                count[type] += 1
                files.append({'path': fullpath, 'type': type})
    
    return files, count

# Function to extract metadata from pdf files.
def get_pdf_metadata(file):
    pdf = PdfReader(file)
    metadata = dict(pdf.metadata)
    metadata['pages'] = len(pdf.pages)
    return metadata

# Function to extract metadata from docx files.
def get_docx_metadata(file):
    doc = docx.Document(file)
    return {
        'author': doc.core_properties.author,
        'category': doc.core_properties.category,
        'comments': doc.core_properties.comments,
        'content_status': doc.core_properties.content_status,
        'created': doc.core_properties.created,
        'identifier': doc.core_properties.identifier,
        'keywords': doc.core_properties.keywords,
        'language': doc.core_properties.language,
        'last_modified_by': doc.core_properties.last_modified_by,
        'last_printed': doc.core_properties.last_printed,
        'modified': doc.core_properties.modified,
        'revision': doc.core_properties.revision,
        'subject': doc.core_properties.subject,
        'title': doc.core_properties.title,
        'version': doc.core_properties.version
    }

# Function to extract metadata from xlsx files.
def get_xlsx_metadata(file):
    wb = load_workbook(file)
    return {
        'creator': wb.properties.creator,
        'title': wb.properties.title,
        'description': wb.properties.description,
        'subject': wb.properties.subject,
        'identifier': wb.properties.identifier,
        'language': wb.properties.language,
        'created': wb.properties.created,
        'modified': wb.properties.modified,
        'lastModifiedBy': wb.properties.lastModifiedBy,
        'category': wb.properties.category,
        'contentStatus': wb.properties.contentStatus,
        'version': wb.properties.version,
        'revision': wb.properties.revision,
        'keywords': wb.properties.keywords,
        'lastPrinted': wb.properties.lastPrinted
    }

# Function to get metadata from a file.
def get_metadata(file):
    if file['type'] == 'pdf':
        return get_pdf_metadata(file['path'])
    if file['type'] == 'docx':
        return get_docx_metadata(file['path'])
    return get_xlsx_metadata(file['path'])

# Function to print metadata.
def print_metadata(file, metadata):
    print("###########################################################################")
    print("File: ", file['path'])
    print("Type: ", file['type'])
    print("Metadata: ")
    for key, value in metadata.items():
        print("\t", key, ":", value)

# Main function.
if __name__ == '__main__':
    print("Using directory: ", args.path)
    files, count = get_files(args.path)
    print("Total files: ", len(files), " PDF: ", count['pdf'], " DOCX: ", count['docx'], " XLSX: ", count['xlsx'])

    # For each file, print the metadata.
    for file in files:
        metadata = get_metadata(file)
        print_metadata(file, metadata)